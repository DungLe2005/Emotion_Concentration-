from copy import deepcopy
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Inches, Pt, RGBColor


ROOT = Path(r"D:\Emotion_Concentration")
TEMPLATE = Path(r"C:\Users\dungl\Downloads\2025.05.27 Mau bia Do an nganh (1).docx")
OUT = ROOT / "Bao_cao_do_an_nganh_EduFocus_AI.docx"


TITLE = (
    "XÂY DỰNG ỨNG DỤNG NHẬN DIỆN CẢM XÚC SINH VIÊN VÀ HỖ TRỢ ĐÁNH GIÁ "
    "MỨC ĐỘ TẬP TRUNG TRONG LỚP HỌC SỬ DỤNG VISION TRANSFORMER"
)


BLACK = RGBColor(0, 0, 0)
BLUE = RGBColor(31, 78, 121)
LIGHT_BLUE = "E8EEF5"
LIGHT_GRAY = "F2F4F7"


def set_run_font(run, name="Times New Roman", size=13, bold=None, italic=None, color=BLACK):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.font.color.rgb = color
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = tc_mar.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            cell.width = widths[idx]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def style_doc(doc):
    section = doc.sections[-1]
    section.top_margin = Cm(2.5)
    section.bottom_margin = Cm(2.5)
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(2.0)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Times New Roman"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.font.size = Pt(13)
    normal.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    normal.paragraph_format.line_spacing = 1.3
    normal.paragraph_format.space_after = Pt(6)

    for name, size, before, after in [
        ("Heading 1", 16, 16, 8),
        ("Heading 2", 14, 12, 6),
        ("Heading 3", 13, 8, 4),
    ]:
        st = styles[name]
        st.font.name = "Times New Roman"
        st._element.rPr.rFonts.set(qn("w:ascii"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:hAnsi"), "Times New Roman")
        st._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
        st.font.size = Pt(size)
        st.font.bold = True
        st.font.color.rgb = BLUE
        st.paragraph_format.space_before = Pt(before)
        st.paragraph_format.space_after = Pt(after)
        st.paragraph_format.keep_with_next = True


def replace_cover_text(doc):
    for p in doc.paragraphs:
        text = p.text.strip()
        if text == "TÊN ĐỀ TÀI":
            p.clear()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(TITLE)
            set_run_font(run, size=16, bold=True)
        elif "GVHD:" in p.text or "SVTH:" in p.text:
            p.clear()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            run = p.add_run("\tGVHD: \t........................................")
            set_run_font(run, size=13, bold=True)
            p.add_run("\n")
            run = p.add_run("\tSVTH: \t........................................")
            set_run_font(run, size=13, bold=True)
            p.add_run("\n")
            run = p.add_run("\tMSSV: \t........................................")
            set_run_font(run, size=13, bold=True)
        elif "Bình Dương, tháng 5 năm 2025" in p.text:
            p.clear()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run("Bình Dương, tháng 7 năm 2026")
            set_run_font(run, size=13, bold=True)


def delete_paragraph(paragraph):
    element = paragraph._element
    element.getparent().remove(element)
    paragraph._p = paragraph._element = None


def trim_cover_to_one_page(doc):
    keep_until = None
    for i, p in enumerate(doc.paragraphs):
        if "GVHD:" in p.text or "SVTH:" in p.text:
            keep_until = i
            break
    if keep_until is None:
        return
    for p in list(doc.paragraphs)[keep_until + 1:]:
        delete_paragraph(p)
    for _ in range(5):
        sp = doc.add_paragraph()
        sp.paragraph_format.space_after = Pt(0)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("Bình Dương, tháng 7 năm 2026")
    set_run_font(run, size=13, bold=True)


def add_page_number(section):
    footer = section.footer
    p = footer.paragraphs[0] if footer.paragraphs else footer.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(0)
    run = p.add_run("Trang ")
    set_run_font(run, size=10)
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = "PAGE"
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "end")
    r = p.add_run()
    r._r.append(fld_char1)
    r._r.append(instr_text)
    r._r.append(fld_char2)


def add_para(doc, text="", style=None, align=None, bold_prefix=None):
    p = doc.add_paragraph(style=style)
    p.paragraph_format.first_line_indent = Cm(1.0) if style is None else None
    p.paragraph_format.line_spacing = 1.3
    p.paragraph_format.space_after = Pt(6)
    if align is not None:
        p.alignment = align
    if bold_prefix and text.startswith(bold_prefix):
        r = p.add_run(bold_prefix)
        set_run_font(r, bold=True)
        r = p.add_run(text[len(bold_prefix):])
        set_run_font(r)
    else:
        r = p.add_run(text)
        set_run_font(r)
    return p


def add_heading(doc, text, level=1):
    p = doc.add_paragraph(style=f"Heading {level}")
    p.paragraph_format.first_line_indent = None
    r = p.add_run(text)
    set_run_font(r, size={1: 16, 2: 14, 3: 13}[level], bold=True, color=BLUE)
    return p


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Paragraph")
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run("- " + item)
        set_run_font(r)


def add_numbered(doc, items):
    for idx, item in enumerate(items, 1):
        p = doc.add_paragraph(style="List Paragraph")
        p.paragraph_format.left_indent = Cm(1.0)
        p.paragraph_format.first_line_indent = Cm(-0.4)
        p.paragraph_format.line_spacing = 1.25
        p.paragraph_format.space_after = Pt(4)
        r = p.add_run(f"{idx}. {item}")
        set_run_font(r)


def add_table(doc, headers, rows, widths):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    set_table_geometry(table, widths)
    for i, header in enumerate(headers):
        cell = table.rows[0].cells[i]
        set_cell_shading(cell, LIGHT_BLUE)
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = p.add_run(header)
        set_run_font(run, size=11, bold=True)
    for row in rows:
        cells = table.add_row().cells
        for i, val in enumerate(row):
            set_cell_margins(cells[i])
            p = cells[i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if i != 0 else WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(val)
            set_run_font(run, size=11)
    doc.add_paragraph()
    return table


def add_toc_field(doc):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("MỤC LỤC")
    set_run_font(r, size=16, bold=True)
    p = doc.add_paragraph()
    fld_char1 = OxmlElement("w:fldChar")
    fld_char1.set(qn("w:fldCharType"), "begin")
    instr_text = OxmlElement("w:instrText")
    instr_text.set(qn("xml:space"), "preserve")
    instr_text.text = 'TOC \\o "1-3" \\h \\z \\u'
    fld_char2 = OxmlElement("w:fldChar")
    fld_char2.set(qn("w:fldCharType"), "separate")
    fld_char3 = OxmlElement("w:fldChar")
    fld_char3.set(qn("w:fldCharType"), "end")
    r = p.add_run()
    r._r.append(fld_char1)
    r._r.append(instr_text)
    r._r.append(fld_char2)
    placeholder = p.add_run("Bấm chuột phải và chọn Update Field để cập nhật mục lục trong Microsoft Word.")
    set_run_font(placeholder, size=11, italic=True, color=RGBColor(90, 90, 90))
    r = p.add_run()
    r._r.append(fld_char3)


def add_report_content(doc):
    style_doc(doc)
    add_page_number(doc.sections[-1])
    doc.add_page_break()
    add_toc_field(doc)
    doc.add_page_break()

    add_heading(doc, "MỞ ĐẦU", 1)
    add_para(
        doc,
        "Trong bối cảnh lớp học trực tuyến và lớp học có hỗ trợ công nghệ ngày càng phổ biến, giảng viên thường thiếu kênh phản hồi tức thời về trạng thái tiếp nhận bài học của sinh viên. Các tín hiệu như biểu cảm khuôn mặt, mức độ tự tin của mô hình nhận diện và phân bố cảm xúc theo thời gian có thể hỗ trợ giảng viên phát hiện sớm thời điểm người học mất tập trung, nhưng việc sử dụng các tín hiệu này cần được triển khai thận trọng, có giới hạn rõ ràng và không thay thế đánh giá sư phạm của con người.",
    )
    add_para(
        doc,
        "Đề tài “Xây dựng ứng dụng nhận diện cảm xúc sinh viên và hỗ trợ đánh giá mức độ tập trung trong lớp học sử dụng Vision Transformer” tập trung xây dựng một hệ thống thử nghiệm có khả năng phát hiện khuôn mặt, chuẩn hóa ảnh mặt, trích xuất đặc trưng thị giác bằng CLIP ViT-B/32, chọn lọc đặc trưng bằng EPSO và phân loại cảm xúc bằng mạng MLP. Kết quả cảm xúc được chuyển thành chỉ số tập trung hỗ trợ quan sát lớp học theo thời gian thực.",
    )
    add_para(
        doc,
        "Phạm vi báo cáo được viết theo hướng khoa học: trình bày vấn đề, khoảng trống nghiên cứu, cơ sở lý thuyết, phương pháp thực nghiệm, kết quả đạt được và so sánh với công trình liên quan. Báo cáo không liệt kê toàn bộ tinh chỉnh nhỏ của pipeline mà tập trung vào các lựa chọn phương pháp có ý nghĩa nghiên cứu.",
    )

    add_heading(doc, "1. TỔNG QUAN ĐỀ TÀI", 1)
    add_heading(doc, "1.1. Bối cảnh và động cơ nghiên cứu", 2)
    add_para(
        doc,
        "Nhận diện cảm xúc khuôn mặt (Facial Expression Recognition - FER) là một hướng quan trọng của thị giác máy tính và tính toán cảm xúc. Trong giáo dục, FER thường được khai thác để suy luận mức độ tham gia, hứng thú, bối rối hoặc mất tập trung của người học. Một số nghiên cứu về lớp học trực tuyến cho thấy camera phổ thông có thể cung cấp các tín hiệu hành vi hữu ích như cảm xúc, chuyển động đầu và mắt để ước lượng trạng thái tham gia học tập [1], [2].",
    )
    add_para(
        doc,
        "Tuy nhiên, bài toán đánh giá tập trung từ khuôn mặt có tính gián tiếp. Một biểu cảm đơn lẻ không đủ để kết luận người học hiểu bài hay không; do đó hệ thống trong đề tài chỉ đóng vai trò hỗ trợ quan sát, hiển thị xác suất cảm xúc và điểm tập trung tham khảo. Quan điểm này phù hợp với định hướng dùng AI như công cụ hỗ trợ ra quyết định, không thay thế vai trò chuyên môn của giảng viên.",
    )

    add_heading(doc, "1.2. Các nghiên cứu liên quan", 2)
    add_para(
        doc,
        "Tài liệu “A Lightweight Teaching Assessment Framework Using Facial Expression Recognition for Online Courses” nhấn mạnh nhu cầu xây dựng khung đánh giá nhẹ cho dạy học trực tuyến dựa trên biểu cảm khuôn mặt [3]. Hướng tiếp cận này gần với mục tiêu của đề tài vì đều khai thác FER để phản hồi trạng thái lớp học. Tuy vậy, các khung nhẹ thường ưu tiên khả năng triển khai và chưa nhấn mạnh việc tận dụng đặc trưng giàu ngữ nghĩa từ mô hình thị giác-ngôn ngữ tiền huấn luyện.",
    )
    add_para(
        doc,
        "Nghiên cứu của Sharma và cộng sự kết hợp cảm xúc, theo dõi mắt và chuyển động đầu để phân loại mức độ tham gia học tập trong kịch bản e-learning [1]. Nair và cộng sự đề xuất phân tích biểu cảm để tăng cường đo lường mức độ tham gia trong học trực tuyến, trong đó cảm xúc phức hợp được xây dựng từ các cảm xúc cơ bản [2]. Các hướng này cho thấy FER có tiềm năng trong giáo dục, nhưng cũng gợi ra nhu cầu hệ thống gọn hơn, dễ triển khai hơn và có thể tận dụng biểu diễn mạnh từ mô hình tiền huấn luyện.",
    )
    add_para(
        doc,
        "Ở tầng mô hình thị giác, Vision Transformer chứng minh rằng ảnh có thể được xử lý như chuỗi patch và đạt hiệu quả cao khi được tiền huấn luyện trên dữ liệu lớn [4]. CLIP mở rộng ý tưởng này bằng học tương phản ảnh-ngôn ngữ trên quy mô lớn, tạo ra encoder ảnh có khả năng chuyển giao sang nhiều bài toán [5]. RetinaFace cung cấp bộ phát hiện và định vị khuôn mặt mạnh trong điều kiện tự nhiên [6]. Các thành phần này là nền tảng cho hướng xây dựng pipeline nhận diện cảm xúc trong đề tài.",
    )

    add_heading(doc, "1.3. Khoảng trống nghiên cứu", 2)
    add_bullets(
        doc,
        [
            "Nhiều hệ thống đánh giá tập trung trong lớp học trực tuyến dựa trên CNN hoặc đặc trưng thủ công, trong khi chưa khai thác đầy đủ biểu diễn ngữ nghĩa giàu thông tin từ các mô hình Vision Transformer tiền huấn luyện.",
            "Một số nghiên cứu dùng nhiều tín hiệu như mắt, đầu, tư thế hoặc tương tác chuột, làm tăng độ phức tạp triển khai; đề tài tập trung vào kênh khuôn mặt để tạo hệ thống nhẹ và dễ tích hợp.",
            "Đặc trưng thị giác kích thước cao giúp mô hình biểu diễn tốt hơn nhưng có thể làm tăng chi phí huấn luyện/suy luận; cần cơ chế chọn lọc đặc trưng để giữ lại thông tin quan trọng.",
            "Các hệ thống FER trong giáo dục thường báo cảm xúc rời rạc; đề tài bổ sung lớp chuyển đổi xác suất cảm xúc thành chỉ số tập trung nhằm hỗ trợ giảng viên quan sát xu hướng lớp học.",
        ],
    )

    add_heading(doc, "1.4. Mục tiêu và đóng góp của đề tài", 2)
    add_numbered(
        doc,
        [
            "Xây dựng pipeline nhận diện cảm xúc khuôn mặt theo thời gian thực gồm phát hiện khuôn mặt, chuẩn hóa ảnh mặt, trích xuất đặc trưng, chọn lọc đặc trưng và phân loại cảm xúc.",
            "Ứng dụng CLIP ViT-B/32 như bộ trích xuất đặc trưng 512 chiều để tận dụng năng lực biểu diễn của Vision Transformer tiền huấn luyện.",
            "Sử dụng EPSO để chọn tập con đặc trưng có ích, giảm nhiễu và hỗ trợ triển khai mô hình phân loại MLP gọn hơn.",
            "Xây dựng ứng dụng web gồm backend FastAPI và frontend React để trình diễn kết quả cảm xúc, xác suất và điểm tập trung.",
        ],
    )

    add_heading(doc, "2. CƠ SỞ LÝ THUYẾT", 1)
    add_heading(doc, "2.1. Nhận diện cảm xúc khuôn mặt", 2)
    add_para(
        doc,
        "FER là bài toán ánh xạ ảnh khuôn mặt sang một trong các nhãn cảm xúc như anger, disgust, fear, happy, neutral, sad, surprise hoặc contempt. Với dữ liệu ảnh tự nhiên, FER chịu ảnh hưởng bởi ánh sáng, góc mặt, chất lượng camera, che khuất và khác biệt cá nhân. Vì vậy, hệ thống cần phát hiện mặt ổn định, chuẩn hóa hình học và dùng bộ trích xuất đặc trưng đủ mạnh trước khi phân loại.",
    )
    add_para(
        doc,
        "Trong đề tài, FER không được xem là kết luận tâm lý tuyệt đối. Đầu ra của mô hình là phân bố xác suất trên các nhãn cảm xúc. Điểm tập trung chỉ là suy luận hỗ trợ từ xác suất này, được thiết kế để cung cấp tín hiệu tham khảo cho giảng viên.",
    )

    add_heading(doc, "2.2. RetinaFace cho phát hiện khuôn mặt", 2)
    add_para(
        doc,
        "RetinaFace là bộ phát hiện khuôn mặt một giai đoạn có khả năng định vị khuôn mặt và landmark trong điều kiện tự nhiên [6]. Landmark mắt được dùng để chuẩn hóa góc nghiêng của khuôn mặt, giúp giảm biến thiên không liên quan đến cảm xúc. Việc tách bước phát hiện mặt ra khỏi bước phân loại giúp pipeline rõ ràng và dễ kiểm soát lỗi.",
    )

    add_heading(doc, "2.3. Vision Transformer và CLIP", 2)
    add_para(
        doc,
        "Vision Transformer chia ảnh thành các patch, biến mỗi patch thành token và xử lý bằng cơ chế self-attention [4]. Cơ chế này cho phép mô hình học quan hệ toàn cục giữa các vùng ảnh, khác với CNN vốn có thiên hướng mạnh về cục bộ. CLIP sử dụng encoder ảnh dựa trên ViT kết hợp học tương phản ảnh-ngôn ngữ, từ đó tạo đặc trưng ảnh có khả năng chuyển giao cao [5].",
    )
    add_para(
        doc,
        "Trong hệ thống EduFocus AI, CLIP ViT-B/32 được dùng như feature extractor. Ảnh mặt sau tiền xử lý được đưa vào encoder ảnh, tạo vector đặc trưng 512 chiều. Thay vì huấn luyện Vision Transformer từ đầu, đề tài tận dụng mô hình tiền huấn luyện để giảm nhu cầu dữ liệu lớn và tăng khả năng tổng quát hóa.",
    )

    add_heading(doc, "2.4. Chọn lọc đặc trưng bằng EPSO", 2)
    add_para(
        doc,
        "Particle Swarm Optimization (PSO) là thuật toán tối ưu hóa quần thể, trong đó mỗi nghiệm ứng viên được cập nhật dựa trên kinh nghiệm cá nhân và nghiệm tốt nhất của quần thể [7]. EPSO trong đề tài được dùng như một biến thể tăng cường cho bài toán chọn đặc trưng: mỗi cá thể biểu diễn một mặt nạ trên vector 512 chiều, hàm mục tiêu kết hợp lỗi phân loại và hệ số phạt khi chọn quá nhiều đặc trưng.",
    )
    add_para(
        doc,
        "Việc chọn lọc đặc trưng có hai mục tiêu: giảm chiều dữ liệu đầu vào cho MLP và loại bỏ thành phần nhiễu hoặc ít liên quan đến biểu cảm. Đây là điểm khác biệt chính so với cách dùng toàn bộ vector đặc trưng tiền huấn luyện mà không kiểm soát độ dư thừa.",
    )

    add_heading(doc, "2.5. MLP, regularization và học theo giáo trình", 2)
    add_para(
        doc,
        "Sau khi giảm chiều, vector đặc trưng được đưa vào mạng MLP gồm các lớp fully connected, batch normalization, LeakyReLU và dropout. CrossEntropyLoss được dùng cho phân loại đa lớp. Mixup là kỹ thuật regularization tạo mẫu huấn luyện bằng tổ hợp lồi giữa hai mẫu và hai nhãn, giúp mô hình bớt ghi nhớ cứng và cải thiện khả năng tổng quát [8].",
    )
    add_para(
        doc,
        "Script huấn luyện trong repo áp dụng warm-up learning rate, cosine annealing warm restarts, gradient clipping và curriculum learning. Về mặt phương pháp, các kỹ thuật này nhằm ổn định tối ưu, giảm overfitting và giúp mô hình học từ dữ liệu sạch trước khi tiếp xúc với nhiễu mạnh hơn.",
    )

    add_heading(doc, "3. THỰC NGHIỆM", 1)
    add_heading(doc, "3.1. Kiến trúc hệ thống", 2)
    add_para(
        doc,
        "Hệ thống được triển khai theo kiến trúc client-server. Backend sử dụng FastAPI để nhận ảnh từ frontend, thực thi pipeline AI và trả về kết quả JSON. Frontend sử dụng React, Vite và TailwindCSS để hiển thị camera/ảnh đầu vào, phân bố cảm xúc, điểm tập trung và log độ trễ từng giai đoạn.",
    )
    add_table(
        doc,
        ["Thành phần", "Công nghệ", "Vai trò"],
        [
            ("Frontend", "React, Vite, TailwindCSS", "Giao diện quan sát kết quả, hiển thị cảm xúc và điểm tập trung."),
            ("Backend", "FastAPI, Python", "Điều phối API, nhận ảnh, gọi pipeline AI và trả kết quả."),
            ("AI core", "PyTorch, TensorFlow, OpenCV", "Chạy CLIP/MLP, RetinaFace và tiền xử lý ảnh."),
            ("Dữ liệu", "labels.csv, Train/Test", "Huấn luyện và đánh giá mô hình nhận diện 8 cảm xúc."),
        ],
        [Cm(3.0), Cm(4.0), Cm(8.0)],
    )

    add_heading(doc, "3.2. Dữ liệu thực nghiệm", 2)
    add_para(
        doc,
        "Dữ liệu nhãn trong workspace gồm 28.175 bản ghi với 8 lớp cảm xúc: anger, contempt, disgust, fear, happy, neutral, sad và surprise. Thư mục ảnh hiện có 16.108 ảnh trong Train và 14.518 ảnh trong Test. Số lượng bản ghi theo nhãn trong labels.csv được tổng hợp ở Bảng 2.",
    )
    add_table(
        doc,
        ["Nhãn", "Số bản ghi"],
        [
            ("surprise", "4.616"),
            ("happy", "4.336"),
            ("anger", "3.608"),
            ("disgust", "3.472"),
            ("contempt", "3.244"),
            ("fear", "3.043"),
            ("sad", "2.995"),
            ("neutral", "2.861"),
        ],
        [Cm(5.0), Cm(5.0)],
    )

    add_heading(doc, "3.3. Quy trình thực nghiệm", 2)
    add_numbered(
        doc,
        [
            "Đọc ảnh và nhãn từ tập dữ liệu, xác định ảnh thuộc Train hoặc Test.",
            "Phát hiện khuôn mặt bằng RetinaFace và chuẩn hóa vùng mặt về kích thước đầu vào của CLIP.",
            "Trích xuất vector đặc trưng 512 chiều bằng CLIP ViT-B/32.",
            "Tối ưu mặt nạ đặc trưng bằng EPSO, dùng KNN và cross-validation 3 fold trong hàm mục tiêu.",
            "Huấn luyện EmotionMLP trên đặc trưng đã chọn; lưu trọng số tốt nhất theo validation accuracy và validation loss.",
            "Đánh giá mô hình bằng accuracy, precision macro, recall macro, F1 macro và confusion matrix.",
            "Tích hợp trọng số vào backend để suy luận theo thời gian thực và trả về cảm xúc, confidence, phân bố xác suất cùng điểm tập trung.",
        ],
    )

    add_heading(doc, "3.4. Công thức hỗ trợ đánh giá tập trung", 2)
    add_para(
        doc,
        "Trong pipeline hiện tại, cảm xúc happy, neutral và surprise được xem là nhóm tín hiệu tích cực hoặc còn chú ý, nên điểm tập trung được tính từ mức nền cao hơn. Các cảm xúc còn lại được xem là tín hiệu cần quan sát thêm, nên điểm nền thấp hơn. Công thức này có tính heuristic và cần được hiệu chỉnh bằng dữ liệu lớp học thực tế trong các nghiên cứu tiếp theo.",
    )
    add_table(
        doc,
        ["Nhóm cảm xúc", "Công thức điểm tập trung", "Ý nghĩa"],
        [
            ("happy, neutral, surprise", "70 + confidence x 30", "Sinh viên có khả năng còn duy trì chú ý hoặc phản hồi tích cực."),
            ("anger, contempt, disgust, fear, sad", "20 + confidence x 30", "Cần theo dõi thêm vì biểu cảm có thể liên quan đến khó chịu, mệt mỏi hoặc mất tập trung."),
        ],
        [Cm(4.2), Cm(4.5), Cm(6.3)],
    )

    add_heading(doc, "4. KẾT QUẢ VÀ ĐÁNH GIÁ", 1)
    add_heading(doc, "4.1. Kết quả triển khai", 2)
    add_para(
        doc,
        "Sản phẩm của đề tài đã hình thành một ứng dụng end-to-end gồm backend suy luận AI và frontend trực quan hóa. Backend có khả năng nạp mô hình RetinaFace, CLIP ViT-B/32 và EmotionMLP; frontend hiển thị cảm xúc chiếm ưu thế, xác suất từng cảm xúc, confidence, ảnh mặt đã xử lý và log độ trễ theo từng bước.",
    )
    add_bullets(
        doc,
        [
            "Pipeline AI được tổ chức thành các giai đoạn rõ ràng: phát hiện khuôn mặt, chuẩn hóa ảnh mặt, trích xuất đặc trưng, chọn lọc đặc trưng, phân loại cảm xúc và tính điểm tập trung.",
            "Mô hình phân loại được đóng gói thành trọng số `emotion_mlp_best_acc_weights.pth` và `emotion_mlp_epso_weights.pth`, cho thấy quá trình huấn luyện đã sinh được artifact phục vụ suy luận.",
            "Dữ liệu nhãn hỗ trợ 8 lớp cảm xúc, mở rộng hơn nhóm 7 cảm xúc phổ biến bằng lớp contempt.",
            "Hệ thống có cơ chế fallback khi chưa có mặt nạ EPSO: dùng toàn bộ 512 chiều để tránh pipeline bị dừng, đồng thời vẫn hỗ trợ nạp mask khi đã huấn luyện xong.",
        ],
    )

    add_heading(doc, "4.2. Kết quả nghiên cứu so với công trình liên quan", 2)
    add_table(
        doc,
        ["Tiêu chí", "Công trình liên quan", "Đề tài thực hiện"],
        [
            ("Nguồn tín hiệu", "Có thể kết hợp cảm xúc, mắt, đầu hoặc hành vi tương tác [1].", "Tập trung vào ảnh khuôn mặt để giảm độ phức tạp triển khai."),
            ("Bộ trích xuất đặc trưng", "Nhiều hướng dùng CNN hoặc mô hình FER chuyên biệt [1], [2].", "Dùng CLIP ViT-B/32 để khai thác biểu diễn tiền huấn luyện 512 chiều."),
            ("Tối ưu đặc trưng", "Không phải mọi hệ thống đều có bước giảm chiều rõ ràng.", "Dùng EPSO chọn đặc trưng nhằm giảm nhiễu và chi phí phân loại."),
            ("Đầu ra sư phạm", "Thường phân loại engagement hoặc cảm xúc.", "Hiển thị cảm xúc, confidence, phân bố xác suất và điểm tập trung tham khảo."),
            ("Triển khai", "Một số nghiên cứu dừng ở mô hình thử nghiệm.", "Có ứng dụng web FastAPI + React phục vụ trình diễn thời gian thực."),
        ],
        [Cm(3.2), Cm(5.9), Cm(5.9)],
    )
    add_para(
        doc,
        "So sánh cho thấy đóng góp chính của đề tài không nằm ở việc đề xuất một mô hình FER hoàn toàn mới, mà ở cách kết hợp các thành phần hiện đại thành pipeline triển khai được: RetinaFace giúp ổn định đầu vào, CLIP ViT-B/32 cung cấp đặc trưng mạnh, EPSO giảm chiều và MLP đảm nhiệm phân loại gọn. Hướng này phù hợp với yêu cầu đồ án ngành vì vừa có cơ sở nghiên cứu, vừa có sản phẩm phần mềm hoàn chỉnh.",
    )

    add_heading(doc, "4.3. Giới hạn và hướng cải tiến", 2)
    add_bullets(
        doc,
        [
            "Điểm tập trung hiện là công thức heuristic dựa trên xác suất cảm xúc, chưa được hiệu chuẩn bằng nhãn tập trung thu thập từ lớp học thực tế.",
            "Dữ liệu FER có thể lệch miền so với bối cảnh lớp học Việt Nam, đặc biệt về ánh sáng, góc camera và biểu cảm tự nhiên.",
            "Hệ thống mới khai thác ảnh khuôn mặt; các tín hiệu bổ sung như hướng nhìn, chuyển động đầu, thời lượng nhìn màn hình hoặc tương tác học tập có thể cải thiện độ tin cậy.",
            "Cần bổ sung báo cáo định lượng sau khi chạy đầy đủ train_epso.py: accuracy, precision, recall, F1, confusion matrix và thời gian suy luận trung bình.",
        ],
    )

    add_heading(doc, "KẾT LUẬN", 1)
    add_para(
        doc,
        "Báo cáo đã trình bày quá trình xây dựng ứng dụng nhận diện cảm xúc sinh viên và hỗ trợ đánh giá mức độ tập trung trong lớp học sử dụng Vision Transformer. Hệ thống tận dụng RetinaFace cho phát hiện khuôn mặt, CLIP ViT-B/32 cho trích xuất đặc trưng, EPSO cho chọn lọc đặc trưng và EmotionMLP cho phân loại 8 cảm xúc. Ứng dụng web giúp trực quan hóa kết quả theo thời gian thực, phù hợp mục tiêu hỗ trợ giảng viên nắm bắt trạng thái lớp học.",
    )
    add_para(
        doc,
        "Về mặt nghiên cứu, đề tài giải quyết khoảng trống giữa mô hình FER học thuật và hệ thống triển khai được trong bối cảnh giáo dục. Việc dùng đặc trưng tiền huấn luyện kết hợp chọn lọc đặc trưng giúp cân bằng giữa chất lượng biểu diễn và tính gọn nhẹ. Trong tương lai, hệ thống cần được đánh giá trên dữ liệu lớp học thật, bổ sung nhãn tập trung do chuyên gia/giảng viên xác nhận và hiệu chỉnh công thức điểm tập trung bằng phương pháp thống kê hoặc học máy.",
    )

    add_heading(doc, "TÀI LIỆU THAM KHẢO", 1)
    refs = [
        '[1] P. Sharma, S. Joshi, S. Gautam, S. Maharjan, S. R. Khanal, M. C. Reis, J. Barroso, and V. M. J. Filipe, "Student Engagement Detection Using Emotion Analysis, Eye Tracking and Head Movement with Machine Learning," arXiv:1909.12913, 2019.',
        '[2] R. R. Nair, T. Babu, and P. K., "Enhancing Student Engagement in Online Learning through Facial Expression Analysis and Complex Emotion Recognition using Deep Learning," arXiv:2311.10343, 2023.',
        '[3] "A Lightweight Teaching Assessment Framework Using Facial Expression Recognition for Online Courses," tài liệu người dùng cung cấp.',
        '[4] A. Dosovitskiy et al., "An Image is Worth 16x16 Words: Transformers for Image Recognition at Scale," in Proc. ICLR, 2021.',
        '[5] A. Radford et al., "Learning Transferable Visual Models From Natural Language Supervision," in Proc. ICML, 2021.',
        '[6] J. Deng, J. Guo, Y. Zhou, J. Yu, I. Kotsia, and S. Zafeiriou, "RetinaFace: Single-stage Dense Face Localisation in the Wild," arXiv:1905.00641, 2019.',
        '[7] J. Kennedy and R. Eberhart, "Particle Swarm Optimization," in Proc. IEEE International Conference on Neural Networks, 1995, pp. 1942-1948.',
        '[8] H. Zhang, M. Cisse, Y. N. Dauphin, and D. Lopez-Paz, "mixup: Beyond Empirical Risk Minimization," in Proc. ICLR, 2018.',
        '[9] I. J. Goodfellow et al., "Challenges in Representation Learning: A Report on Three Machine Learning Contests," in Proc. ICONIP, 2013.',
    ]
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.first_line_indent = Cm(-0.7)
        p.paragraph_format.left_indent = Cm(0.7)
        p.paragraph_format.line_spacing = 1.15
        p.paragraph_format.space_after = Pt(5)
        r = p.add_run(ref)
        set_run_font(r, size=11)


def main():
    doc = Document(str(TEMPLATE))
    replace_cover_text(doc)
    trim_cover_to_one_page(doc)
    add_report_content(doc)
    doc.save(str(OUT))
    print(OUT)


if __name__ == "__main__":
    main()
